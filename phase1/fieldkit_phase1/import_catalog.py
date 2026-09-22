#!/usr/bin/env python3
"""
FieldKit: Catalog import from a ServiceFusion "Company Services" export.
Directive Stage 4, Increment 5.6.

Reads a ServiceFusion CompanyServices_*.xlsx export (header row 1: Service
Category, Service Name, Description, Status, Regular Rate, Cost, Member Rate,
Add-On Rate, Premium Rate, Value Rate, Afterhours Rate) and inserts new
catalog_items rows for the real, billable service categories only.

Scope decided with Chris (2026-09-22), reviewing this exact file for Get a
Grip:
  - Only rows where Status == 'Active' AND category in REAL_CATEGORIES below.
  - 'QBO Service' and 'Service' categories are excluded outright -- they mix
    ServiceFusion/QuickBooks bookkeeping rows (Labor, Drive Time) and, more
    importantly, TAX LINE ITEMS disguised as services ("County Tax 2%",
    "Transit Tax") that would double up against FieldKit's own tax engine if
    imported as catalog items. A few of these rows do carry a real price
    (Min Charge $185, OCC FEE $45, Xtra Prep/Full Vinyl/Post Clean $45) --
    deliberately NOT auto-imported since Chris's review didn't specifically
    resolve those three; add them by hand via /settings/catalog/new if
    wanted, it's a two-minute job.
  - 'Discount' category is excluded -- a discount isn't a catalog item in
    FieldKit's model (see invoice_adjustments, adjustment_type='discount').
  - Rows with no Regular Rate in a real category (most Resurfacing/Repairs/
    Stripping rows) are quoted per job in ServiceFusion, not fixed-priced --
    Chris confirmed (2026-09-22) these still need a catalog entry (so they're
    selectable on a WO/invoice) with DEFAULT_UNPRICED_RATE ($50, "the
    starting price") as a deliberately-visible placeholder, not a real rate.
  - unit_of_measure has no source column at all -- defaulted to 'each',
    matching the one real overlapping precedent already in FieldKit
    (Bathtub Resurface, $275, unit_of_measure='each').
  - is_taxable has no source column either -- defaulted to True, same
    precedent (Bathtub Resurface is_taxable=TRUE). NC taxes contractor
    repair/resurfacing labor; confirm/adjust per item once real invoices
    start flowing if any of these turn out exempt.

Usage:
  Dry run (default):
    docker compose exec -T app python phase1/fieldkit_phase1/import_catalog.py \\
        /path/to/CompanyServices_*.xlsx getagrip

  Real import:
    docker compose exec -T app python phase1/fieldkit_phase1/import_catalog.py \\
        /path/to/CompanyServices_*.xlsx getagrip --commit
"""
import sys
import os
import argparse
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, '/app')
from app import get_db_connection  # noqa: E402

REAL_CATEGORIES = {'Cleaning', 'Resurfacing', 'Repairs', 'Stripping', 'Tile Work'}
DEFAULT_UNPRICED_RATE = 50.00
DEFAULT_UNIT_OF_MEASURE = 'each'
DEFAULT_IS_TAXABLE = True


def read_export(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    items = []
    for r in rows:
        cat, name, desc, status, reg, cost, member, addon, premium, value, afterhours = r
        if status != 'Active' or cat not in REAL_CATEGORIES:
            continue
        price = float(reg) if reg is not None else DEFAULT_UNPRICED_RATE
        items.append({
            'name': (name or '').strip(),
            'category': cat,
            'description': (desc or '').strip() or None,
            'unit_price': price,
            'was_unpriced': reg is None,
        })
    return items


def run(xlsx_path, company_key, dry_run, username):
    items = read_export(xlsx_path)
    conn = get_db_connection(company_key)
    cur = conn.cursor()
    cur.execute("SELECT lower(name) AS n FROM catalog_items WHERE deleted_at IS NULL")
    existing_names = {r['n'] for r in cur.fetchall()}

    inserted = skipped_existing = 0
    unpriced_count = 0
    for item in items:
        if item['name'].lower() in existing_names:
            skipped_existing += 1
            continue
        if item['was_unpriced']:
            unpriced_count += 1
        print(f"  {'[DRY RUN] would insert' if dry_run else '[INSERT]'} "
              f"{item['category']:12} | {item['name']:35} | ${item['unit_price']:.2f}"
              f"{'  (placeholder price -- quoted per job)' if item['was_unpriced'] else ''}")
        if not dry_run:
            cur.execute("""
                INSERT INTO catalog_items
                    (billing_behavior, name, default_description, category,
                     unit_price, unit_of_measure, is_taxable,
                     sort_order, is_active, created_by, updated_by)
                VALUES ('standard', %s, %s, %s, %s, %s, %s,
                    (SELECT COALESCE(MAX(sort_order),0)+1 FROM catalog_items WHERE deleted_at IS NULL),
                    TRUE, %s, %s)
            """, (item['name'], item['description'], item['category'],
                  item['unit_price'], DEFAULT_UNIT_OF_MEASURE, DEFAULT_IS_TAXABLE,
                  username, username))
        inserted += 1

    if not dry_run:
        conn.commit()
    cur.close(); conn.close()

    print(f"\n{company_key}: {inserted} catalog item(s) {'would be inserted' if dry_run else 'inserted'} "
          f"({unpriced_count} at the ${DEFAULT_UNPRICED_RATE:.0f} placeholder rate), "
          f"{skipped_existing} already existed (skipped, name match).")
    if dry_run:
        print("DRY RUN -- nothing was written. Re-run with --commit to apply.")


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('xlsx_path', help='Path to the CompanyServices_*.xlsx export')
    parser.add_argument('company_key', choices=['getagrip', 'kleanit_charlotte', 'cts', 'kleanit_sf'])
    parser.add_argument('--commit', action='store_true', help='Actually write. Default is a dry run.')
    parser.add_argument('--username', default='import_catalog.py')
    args = parser.parse_args()
    run(args.xlsx_path, args.company_key, dry_run=not args.commit, username=args.username)
