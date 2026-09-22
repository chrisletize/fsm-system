#!/usr/bin/env python3
"""
FieldKit: customer_job_dates import direct from ServiceFusion weekly
"Customer Revenue Report" exports. Supplements (does not replace)
import_job_dates.py, which sources the same table from the sibling
statements database -- that source has zero job history for three of the
four companies (D-038) and, for Get a Grip, only reaches as far forward as
whatever was last loaded there. These direct SF exports are more current.

Source format: one file per week, header row 1, read_only=True required
(openpyxl chokes on this export's embedded comments otherwise). Relevant
columns: Customer, Date, Status.

Scope decided with Chris (2026-09-22): only rows with Status in
DONE_STATUSES count as a real service date -- 'Invoiced'/'Completed', the
same two statuses the recency report's own WO-derived formula counts
(D-084). 'Scheduled' (not done yet), 'Dispatched', and 'Started' are
excluded -- a job that's merely en route or in progress isn't "service
performed" for recency purposes.

Accepts a directory of weekly files (all *.xlsx in it are read) rather than
one file at a time, since this data only makes sense in aggregate.

Usage:
  Dry run (default):
    docker compose exec -T app python phase1/fieldkit_phase1/import_job_history_direct.py \\
        /path/to/job-history-dir getagrip

  Real import:
    docker compose exec -T app python phase1/fieldkit_phase1/import_job_history_direct.py \\
        /path/to/job-history-dir getagrip --commit
"""
import sys
import os
import re
import csv
import glob
import argparse
from datetime import datetime
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, '/app')
from app import get_db_connection  # noqa: E402

DONE_STATUSES = {'Invoiced', 'Completed'}


def normalize_name(name):
    if not name:
        return ''
    s = name.strip().lower()
    s = re.sub(r"[.,'\"]", '', s)
    s = re.sub(r'\s+', ' ', s)
    return s


def read_export_dir(dir_path):
    by_customer = {}
    skipped_status = 0
    for fp in sorted(glob.glob(os.path.join(dir_path, '*.xlsx'))):
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            customer, date_str, status = row[0], row[8], row[10]
            if status not in DONE_STATUSES:
                skipped_status += 1
                continue
            if not customer or not date_str:
                continue
            d = datetime.strptime(date_str, '%m/%d/%Y').date()
            by_customer.setdefault(customer, set()).add(d)
    return by_customer, skipped_status


def run(dir_path, company_key, dry_run, username, out_dir):
    by_customer, skipped_status = read_export_dir(dir_path)
    total_dates = sum(len(v) for v in by_customer.values())
    print(f"Source: {len(by_customer)} distinct customer(s), {total_dates} distinct service date(s) "
          f"({skipped_status} row(s) skipped -- not Invoiced/Completed).\n")

    conn = get_db_connection(company_key)
    cur = conn.cursor()
    cur.execute("SELECT id, property_name FROM customers WHERE deleted_at IS NULL")
    by_name = {}
    for r in cur.fetchall():
        by_name.setdefault(normalize_name(r['property_name']), []).append(r['id'])

    matched_customers = inserted = already_present = 0
    unmatched = {}

    for sf_name, dates in by_customer.items():
        ids = by_name.get(normalize_name(sf_name))
        if not ids:
            unmatched[sf_name] = len(dates)
            continue
        if len(ids) > 1:
            print(f"  [SKIP] '{sf_name}' matches {len(ids)} customers -- ambiguous, needs a human.")
            continue
        matched_customers += 1
        customer_id = ids[0]
        for d in dates:
            if dry_run:
                cur.execute("SELECT 1 FROM customer_job_dates WHERE customer_id = %s AND job_date = %s",
                            (customer_id, d))
                if cur.fetchone():
                    already_present += 1
                else:
                    inserted += 1
            else:
                cur.execute("""
                    INSERT INTO customer_job_dates (customer_id, job_date, source, created_by)
                    VALUES (%s, %s, 'servicefusion_import_direct', %s)
                    ON CONFLICT (customer_id, job_date) DO NOTHING
                    RETURNING id
                """, (customer_id, d, username))
                if cur.fetchone():
                    inserted += 1
                else:
                    already_present += 1

    if not dry_run:
        conn.commit()
    cur.close(); conn.close()

    print(f"{company_key}: {matched_customers} customer(s) matched, "
          f"{inserted} job date row(s) {'would be inserted' if dry_run else 'inserted'}, "
          f"{already_present} already present (e.g. from a real FieldKit WO on the same day), "
          f"{len(unmatched)} unmatched customer name(s).")

    if unmatched:
        path = os.path.join(out_dir, f'unmatched_job_history_direct_{company_key}.csv')
        with open(path, 'w', newline='') as f:
            w = csv.writer(f)
            w.writerow(['servicefusion_customer_name', 'date_count'])
            for name, count in sorted(unmatched.items(), key=lambda x: -x[1]):
                w.writerow([name, count])
        print(f"Unmatched names logged to {path}")

    if dry_run:
        print("\nDRY RUN -- nothing was written. Re-run with --commit to apply.")


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('dir_path', help='Directory containing the weekly CustomerRevenueReport_*.xlsx files')
    parser.add_argument('company_key', choices=['getagrip', 'kleanit_charlotte', 'cts', 'kleanit_sf'])
    parser.add_argument('--commit', action='store_true', help='Actually write. Default is a dry run.')
    parser.add_argument('--username', default='import_job_history_direct.py')
    parser.add_argument('--out-dir', default='/tmp')
    args = parser.parse_args()
    run(args.dir_path, args.company_key, dry_run=not args.commit, username=args.username, out_dir=args.out_dir)
